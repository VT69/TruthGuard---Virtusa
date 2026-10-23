"""Render the Northwind Bank synthetic corpus, its manifest and the gold question set.

Deterministic and seeded: `python eval/generator/generate.py` always produces identical
files. Documents are written as Markdown, DOCX, PDF, HTML and plain text so every parser
path is exercised; metadata reaches the parser in four styles (see corpus.yaml).

Outputs:
  data/sample_corpus/<files>          the corpus
  data/sample_corpus/manifest.json    ground truth: facts, lineage edges, planted conflicts
  eval/datasets/gold.jsonl            all gold questions with their split
  eval/datasets/{dev,test}.jsonl      70/30 split, stratified by case; test is frozen
"""

from __future__ import annotations

import argparse
import html
import json
import random
import re
from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Any

import yaml

ROOT = Path(__file__).resolve().parents[2]
GEN_DIR = Path(__file__).resolve().parent
CORPUS_DIR = ROOT / "data" / "sample_corpus"
DATASET_DIR = ROOT / "eval" / "datasets"
SEED = 7
STATUS_RANK = {
    "VERIFIED": 0,
    "RESOLVED_CONFLICT": 1,
    "UNRESOLVED_CONFLICT": 2,
    "INSUFFICIENT_EVIDENCE": -1,
}
DOC_TYPE_LABEL = {
    "policy": "Policy",
    "board_policy": "Board Policy",
    "circular": "Circular",
    "procedure": "Procedure",
    "amendment": "Amendment",
    "addendum": "Addendum",
    "product_terms": "Product Terms",
    "schedule": "Schedule of Charges",
    "notice": "Notice",
    "faq": "FAQ",
    "help_article": "Help Article",
    "memo": "Memo",
    "email": "Email",
}


def long_date(d: date) -> str:
    return f"{d.day} {d.strftime('%B %Y')}"


# --------------------------------------------------------------------------- body model


def parse_body(body: str) -> list[tuple[str, int, str]]:
    """Tiny markdown subset -> blocks of (kind, level, text); kind in heading|bullet|para."""
    blocks: list[tuple[str, int, str]] = []
    for line in body.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if m:
            blocks.append(("heading", len(m.group(1)), m.group(2)))
        elif stripped.startswith("- "):
            blocks.append(("bullet", 0, stripped[2:]))
        else:
            blocks.append(("para", 0, stripped))
    return blocks


def header_lines(doc: dict[str, Any]) -> list[str]:
    meta = doc["meta"]
    lines: list[str] = []
    if meta["doc_type"] in ("email", "memo"):
        lines.append(f"From: {meta.get('issuer', 'Northwind Bank')}")
        if meta.get("published"):
            lines.append(f"Date: {long_date(meta['published'])}")
        return lines
    parts = [f"Document type: {DOC_TYPE_LABEL[meta['doc_type']]}"]
    if meta.get("version"):
        parts.append(f"Version: {meta['version']}")
    lines.append(" | ".join(parts))
    if meta.get("effective_from"):
        lines.append(f"Effective from: {long_date(meta['effective_from'])}")
    if meta.get("published"):
        lines.append(f"Issued on: {long_date(meta['published'])}")
    if meta.get("issuer"):
        lines.append(f"Issued by: {meta['issuer']}")
    return lines


def explicit_meta(doc: dict[str, Any]) -> dict[str, Any]:
    meta = doc["meta"]
    out: dict[str, Any] = {
        "doc_id": doc["key"],
        "title": meta["title"],
        "doc_type": meta["doc_type"],
    }
    for k in ("version", "effective_from", "published", "topic", "series", "issuer"):
        if meta.get(k) is not None:
            out[k] = meta[k].isoformat() if isinstance(meta[k], date) else meta[k]
    return out


# --------------------------------------------------------------------------- renderers


def render_md(doc: dict[str, Any], path: Path) -> None:
    text = doc["body"].rstrip() + "\n"
    if doc["meta_style"] == "frontmatter":
        fm = yaml.safe_dump(explicit_meta(doc), sort_keys=False, allow_unicode=True)
        text = f"---\n{fm}---\n\n{text}"
    elif doc["meta_style"] == "header":
        blocks = text.split("\n", 1)
        text = (
            blocks[0] + "\n\n" + "\n".join(header_lines(doc)) + "\n" + (blocks[1] if len(blocks) > 1 else "")
        )
    path.write_text(text, encoding="utf-8")


def render_txt(doc: dict[str, Any], path: Path) -> None:
    lines = doc["body"].rstrip().split("\n")
    if doc["meta_style"] == "header":
        lines = [lines[0], "", *header_lines(doc), *lines[1:]]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def render_html(doc: dict[str, Any], path: Path) -> None:
    meta_tags = ""
    if doc["meta_style"] == "properties":
        meta_tags = "\n".join(
            f'    <meta name="{k}" content="{html.escape(str(v))}">' for k, v in explicit_meta(doc).items()
        )
    out = [
        "<!doctype html>",
        '<html lang="en">',
        "  <head>",
        '    <meta charset="utf-8">',
        f"    <title>{html.escape(doc['meta']['title'])}</title>",
    ]
    if meta_tags:
        out.append(meta_tags)
    out += ["  </head>", "  <body>"]
    in_list = False
    for kind, level, text in parse_body(doc["body"]):
        if kind == "bullet" and not in_list:
            out.append("    <ul>")
            in_list = True
        if kind != "bullet" and in_list:
            out.append("    </ul>")
            in_list = False
        if kind == "heading":
            out.append(f"    <h{level}>{html.escape(text)}</h{level}>")
        elif kind == "bullet":
            out.append(f"      <li>{html.escape(text)}</li>")
        else:
            out.append(f"    <p>{html.escape(text)}</p>")
    if in_list:
        out.append("    </ul>")
    out += ["  </body>", "</html>"]
    path.write_text("\n".join(out) + "\n", encoding="utf-8")


def render_docx(doc: dict[str, Any], path: Path) -> None:
    from docx import Document as DocxDocument

    d = DocxDocument()
    meta = doc["meta"]
    if doc["meta_style"] == "properties":
        cp = d.core_properties
        cp.title = meta["title"]
        cp.subject = meta["doc_type"]
        cp.category = meta.get("topic") or ""
        cp.keywords = meta.get("series") or ""
        cp.version = str(meta.get("version") or "")
        cp.author = meta.get("issuer") or ""
        cp.identifier = doc["key"]
        extra = []
        for k in ("effective_from", "published"):
            if meta.get(k):
                extra.append(f"{k}={meta[k].isoformat()}")
        cp.comments = "; ".join(extra)
    for kind, level, text in parse_body(doc["body"]):
        if kind == "heading":
            d.add_heading(text, level=level)
        elif kind == "bullet":
            d.add_paragraph(text, style="List Bullet")
        else:
            d.add_paragraph(text)
    d.save(str(path))
    _normalize_docx_timestamps(path)


def _normalize_docx_timestamps(path: Path) -> None:
    """python-docx stamps created/modified with 'now'; pin them so output is reproducible."""
    import zipfile

    data: dict[str, bytes] = {}
    with zipfile.ZipFile(path) as zf:
        for name in zf.namelist():
            data[name] = zf.read(name)
    core = data["docProps/core.xml"].decode("utf-8")
    core = re.sub(
        r"(<dcterms:(?:created|modified)[^>]*>)[^<]*(<)",
        r"\g<1>2020-01-01T00:00:00Z\2",
        core,
    )
    data["docProps/core.xml"] = core.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name in sorted(data):
            info = zipfile.ZipInfo(name, date_time=(2020, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            zf.writestr(info, data[name])


FONT_CANDIDATES = [
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("/usr/share/fonts/TTF/DejaVuSans.ttf"),
    Path("/Library/Fonts/Arial Unicode.ttf"),
    Path("C:/Windows/Fonts/arial.ttf"),
]


def render_pdf(doc: dict[str, Any], path: Path) -> None:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_creation_date(__import__("datetime").datetime(2020, 1, 1))
    pdf.set_title(doc["meta"]["title"])
    pdf.add_page()
    font = next((f for f in FONT_CANDIDATES if f.exists()), None)
    family = "Body"
    if font:
        pdf.add_font(family, "", str(font))
        pdf.add_font(family, "B", str(font))
    else:  # core font fallback cannot draw the rupee sign
        family = "Helvetica"

    def clean(text: str) -> str:
        return text if font else text.replace("₹", "Rs. ")

    blocks = parse_body(doc["body"])
    title = blocks[0][2]
    pdf.set_font(family, "B", 15)
    pdf.multi_cell(0, 8, clean(title), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(family, "", 10)
    if doc["meta_style"] == "header":
        for line in header_lines(doc):
            pdf.multi_cell(0, 6, clean(line), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    for kind, level, text in blocks[1:]:
        if kind == "heading":
            pdf.set_font(family, "B", 13 if level == 2 else 11)
            pdf.ln(2)
            pdf.multi_cell(0, 7, clean(text), new_x="LMARGIN", new_y="NEXT")
            pdf.set_font(family, "", 10)
        elif kind == "bullet":
            pdf.multi_cell(0, 6, clean("- " + text), new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.multi_cell(0, 6, clean(text), new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))


RENDERERS = {
    "md": render_md,
    "txt": render_txt,
    "html": render_html,
    "docx": render_docx,
    "pdf": render_pdf,
}


# --------------------------------------------------------------------------- manifest


def build_manifest(corpus: dict[str, Any]) -> dict[str, Any]:
    docs = []
    for doc in corpus["documents"]:
        meta = {k: (v.isoformat() if isinstance(v, date) else v) for k, v in doc["meta"].items()}
        docs.append(
            {
                "key": doc["key"],
                "filename": doc["filename"],
                "format": doc["format"],
                "meta_style": doc["meta_style"],
                "meta": meta,
                "injection": bool(doc.get("injection", False)),
                "facts": [
                    {"subject": s, "attribute": a, "scope": sc, "value": v} for s, a, sc, v in doc["facts"]
                ],
            }
        )
    conflicts = []
    for pc in corpus["planted_conflicts"]:
        item = {k: (v.isoformat() if isinstance(v, date) else v) for k, v in pc.items()}
        item.setdefault("as_of", corpus["eval_today"].isoformat())
        conflicts.append(item)
    return {
        "name": "northwind-bank",
        "eval_today": corpus["eval_today"].isoformat(),
        "documents": docs,
        "edges": corpus["edges"],
        "planted_conflicts": conflicts,
        "numeric_drift_controls": corpus["numeric_drift_controls"],
    }


# --------------------------------------------------------------------------- questions


def _expectation_record(exp: dict[str, Any], default_as_of: date) -> dict[str, Any]:
    as_of = exp.get("as_of", default_as_of)
    status = exp["status"]
    winners = exp.get("winners", [])
    positions = exp.get("positions", [])
    return {
        "expectation": exp["id"],
        "case": exp["case"],
        "slots": [exp["slot"]],
        "scope": exp.get("scope"),
        "as_of": as_of.isoformat() if isinstance(as_of, date) else as_of,
        "expected_status": status,
        "expected_values": exp["values"],
        "expected_winners": winners,
        "expected_eliminations": exp.get("eliminated", {}),
        "expected_positions": positions,
        "expected_flags": exp.get("flags", []),
        "injection_docs": exp.get("injection_docs", []),
        "must_cite": positions
        if status == "UNRESOLVED_CONFLICT"
        else winners[:1]
        if len(winners) == 1
        else winners,
        "must_disclose_conflict": status in ("RESOLVED_CONFLICT", "UNRESOLVED_CONFLICT"),
    }


def build_questions(qspec: dict[str, Any]) -> list[dict[str, Any]]:
    default_as_of: date = qspec["defaults"]["as_of"]
    by_id = {e["id"]: e for e in qspec["expectations"]}
    out: list[dict[str, Any]] = []
    for exp in qspec["expectations"]:
        base = _expectation_record(exp, default_as_of)
        kind = "as_of" if "as_of" in exp else ("control" if exp["status"] == "VERIFIED" else "planted")
        for phr in exp["phrasings"]:
            out.append({**base, "question": phr, "kind": kind, "claimed_value": None})
        for bait in exp.get("bait", []):
            out.append({**base, "question": bait, "kind": "sycophancy", "claimed_value": True})
    for mh in qspec["multi_hop"]:
        parts = [_expectation_record(by_id[p], default_as_of) for p in mh["parts"]]
        status = max((p["expected_status"] for p in parts), key=lambda s: STATUS_RANK[s])
        merged: dict[str, Any] = {
            "expectation": "+".join(mh["parts"]),
            "case": "multi_hop",
            "slots": [s for p in parts for s in p["slots"]],
            "scope": None,
            "as_of": default_as_of.isoformat(),
            "expected_status": status,
            "expected_values": [v for p in parts for v in p["expected_values"]],
            "expected_winners": [w for p in parts for w in p["expected_winners"]],
            "expected_eliminations": {k: v for p in parts for k, v in p["expected_eliminations"].items()},
            "expected_positions": [d for p in parts for d in p["expected_positions"]],
            "expected_flags": sorted({f for p in parts for f in p["expected_flags"]}),
            "injection_docs": sorted({d for p in parts for d in p["injection_docs"]}),
            "must_cite": [d for p in parts for d in p["must_cite"]],
            "must_disclose_conflict": any(p["must_disclose_conflict"] for p in parts),
            "question": mh["question"],
            "kind": "multi_hop",
            "claimed_value": None,
        }
        out.append(merged)
    for q in qspec["unanswerable"]:
        out.append(
            {
                "expectation": "unanswerable",
                "case": "unanswerable",
                "slots": [],
                "scope": None,
                "as_of": default_as_of.isoformat(),
                "expected_status": "INSUFFICIENT_EVIDENCE",
                "expected_values": [],
                "expected_winners": [],
                "expected_eliminations": {},
                "expected_positions": [],
                "expected_flags": [],
                "injection_docs": [],
                "must_cite": [],
                "must_disclose_conflict": False,
                "question": q,
                "kind": "unanswerable",
                "claimed_value": None,
            }
        )
    rng = random.Random(SEED)
    by_case: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for q in out:
        by_case[q["case"]].append(q)
    for case in sorted(by_case):
        items = by_case[case]
        rng.shuffle(items)
        n_test = max(1, round(len(items) * 0.3)) if len(items) > 1 else 0
        for i, q in enumerate(items):
            q["split"] = "test" if i < n_test else "dev"
    for i, q in enumerate(out, start=1):
        q["id"] = f"q{i:03d}"
    return out


# --------------------------------------------------------------------------- main


def validate(corpus: dict[str, Any]) -> list[str]:
    """Cheap self-checks so a typo in corpus.yaml cannot silently corrupt ground truth."""
    problems: list[str] = []
    keys = {d["key"] for d in corpus["documents"]}
    for e in corpus["edges"]:
        for end in ("from", "to"):
            if e[end] not in keys:
                problems.append(f"edge references unknown doc {e[end]}")
    for pc in corpus["planted_conflicts"]:
        for k in ("winner", "loser"):
            if pc.get(k) and pc[k] not in keys:
                problems.append(f"{pc['id']}: unknown {k} {pc[k]}")
    return problems


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=CORPUS_DIR)
    parser.add_argument("--datasets", type=Path, default=DATASET_DIR)
    args = parser.parse_args()

    corpus = yaml.safe_load((GEN_DIR / "corpus.yaml").read_text(encoding="utf-8"))
    qspec = yaml.safe_load((GEN_DIR / "questions.yaml").read_text(encoding="utf-8"))
    problems = validate(corpus)
    if problems:
        raise SystemExit("corpus.yaml problems:\n" + "\n".join(problems))

    args.out.mkdir(parents=True, exist_ok=True)
    for old in args.out.iterdir():
        if old.is_file():
            old.unlink()
    for doc in corpus["documents"]:
        RENDERERS[doc["format"]](doc, args.out / doc["filename"])
    manifest = build_manifest(corpus)
    (args.out / "manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", "utf-8"
    )

    questions = build_questions(qspec)
    args.datasets.mkdir(parents=True, exist_ok=True)
    for name, rows in (
        ("gold", questions),
        ("dev", [q for q in questions if q["split"] == "dev"]),
        ("test", [q for q in questions if q["split"] == "test"]),
    ):
        with (args.datasets / f"{name}.jsonl").open("w", encoding="utf-8") as fh:
            for row in rows:
                fh.write(json.dumps(row, ensure_ascii=False) + "\n")

    n_conf = len(manifest["planted_conflicts"])
    cases = sorted({pc["case"] for pc in manifest["planted_conflicts"]})
    print(
        f"documents: {len(manifest['documents'])}  planted conflicts: {n_conf}  edges: {len(manifest['edges'])}"
    )
    print(f"conflict cases: {', '.join(cases)}")
    by_split = defaultdict(int)
    for q in questions:
        by_split[q["split"]] += 1
    print(f"questions: {len(questions)}  dev: {by_split['dev']}  test: {by_split['test']}")


if __name__ == "__main__":
    main()
