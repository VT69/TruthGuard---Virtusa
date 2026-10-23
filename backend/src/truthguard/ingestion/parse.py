"""Parse Markdown, text, HTML, DOCX and PDF into sections with heading paths and offsets.

Every format is reduced to a list of blocks (heading / bullet / paragraph) and rendered to
one canonical `raw_text`; section offsets index into that text, so cited quotes can be
verified verbatim later. Docling is the documented upgrade path for complex layouts
(ADR 0004); these light parsers keep the default install small and CPU-only.
"""

from __future__ import annotations

import io
import re
from collections import Counter
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

import yaml

from truthguard.core.errors import IngestError
from truthguard.domain.models import ParsedDocument, Section

KV_LINE = re.compile(r"^[A-Z][A-Za-z ]{1,30}:\s")


@dataclass
class Block:
    kind: str  # heading | bullet | para
    text: str
    level: int = 0


def blocks_to_document(
    blocks: list[Block],
    fmt: str,
    frontmatter: dict[str, Any] | None = None,
    properties: dict[str, Any] | None = None,
) -> ParsedDocument:
    lines: list[str] = []
    offsets: list[tuple[Block, int, int]] = []
    pos = 0
    for b in blocks:
        text = b.text.strip()
        if not text:
            continue
        rendered = f"- {text}" if b.kind == "bullet" else text
        if b.kind == "heading" and lines:
            lines.append("")
            pos += 1
        offsets.append((b, pos, pos + len(rendered)))
        lines.append(rendered)
        pos += len(rendered) + 1
    raw_text = "\n".join(lines)

    sections: list[Section] = []
    stack: list[tuple[int, str]] = []
    current_start: int | None = None
    current_path: list[str] = []

    def close(end: int) -> None:
        if current_start is None:
            return
        body = raw_text[current_start:end].rstrip()
        if body.strip():
            sections.append(
                Section(
                    heading_path=list(current_path),
                    text=body,
                    char_start=current_start,
                    char_end=current_start + len(body),
                )
            )

    for b, start, _end in offsets:
        if b.kind == "heading":
            close(start)
            while stack and stack[-1][0] >= b.level:
                stack.pop()
            stack.append((b.level, b.text.strip()))
            current_path = [h for _, h in stack]
            current_start = start
        elif current_start is None:
            current_start = start
            current_path = []
    close(len(raw_text))
    return ParsedDocument(
        raw_text=raw_text,
        sections=sections,
        frontmatter=frontmatter or {},
        properties=properties or {},
        fmt=fmt,
    )


# --------------------------------------------------------------------------- markdown / text


def split_frontmatter(text: str) -> tuple[dict[str, Any], str]:
    if text.startswith("---\n"):
        end = text.find("\n---", 4)
        if end != -1:
            data = yaml.safe_load(text[4:end]) or {}
            if isinstance(data, dict):
                return {str(k): v for k, v in data.items()}, text[end + 4 :].lstrip("\n")
    return {}, text


def parse_markdown(text: str) -> ParsedDocument:
    fm, body = split_frontmatter(text)
    blocks: list[Block] = []
    for line in body.splitlines():
        s = line.strip()
        if not s:
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            blocks.append(Block("heading", m.group(2), len(m.group(1))))
        elif re.match(r"^[-*+]\s+", s):
            blocks.append(Block("bullet", re.sub(r"^[-*+]\s+", "", s)))
        else:
            blocks.append(Block("para", s))
    return blocks_to_document(blocks, "md", frontmatter=fm)


def parse_text(text: str) -> ParsedDocument:
    blocks: list[Block] = []
    first = True
    for line in text.splitlines():
        s = line.strip()
        if not s:
            continue
        if first:
            blocks.append(Block("heading", s, 1))
            first = False
        elif s.endswith("?") and len(s) < 160 and not KV_LINE.match(s):
            blocks.append(Block("heading", s, 2))
        elif re.match(r"^[-*•]\s+", s):
            blocks.append(Block("bullet", re.sub(r"^[-*•]\s+", "", s)))
        else:
            blocks.append(Block("para", s))
    return blocks_to_document(blocks, "txt")


# --------------------------------------------------------------------------- html


class _HTMLCollector(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[Block] = []
        self.meta: dict[str, str] = {}
        self.title = ""
        self._tag: str | None = None
        self._buf: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        a = {k: v or "" for k, v in attrs}
        if tag in ("script", "style"):
            self._skip += 1
        elif tag == "meta" and "name" in a:
            self.meta[a["name"].lower()] = a.get("content", "")
        elif tag in ("h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "title", "td", "dt", "dd"):
            self._flush()
            self._tag = tag
        elif tag == "br":
            self._buf.append(" ")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style"):
            self._skip = max(0, self._skip - 1)
        elif tag == self._tag:
            self._flush()

    def handle_data(self, data: str) -> None:
        if not self._skip and self._tag:
            self._buf.append(data)

    def _flush(self) -> None:
        text = re.sub(r"\s+", " ", "".join(self._buf)).strip()
        tag, self._tag, self._buf = self._tag, None, []
        if not text or tag is None:
            return
        if tag == "title":
            self.title = text
        elif tag.startswith("h"):
            self.blocks.append(Block("heading", text, int(tag[1])))
        elif tag == "li":
            self.blocks.append(Block("bullet", text))
        else:
            self.blocks.append(Block("para", text))


def parse_html(text: str) -> ParsedDocument:
    c = _HTMLCollector()
    c.feed(text)
    c.close()
    c._flush()
    blocks = c.blocks
    if c.title and not any(b.kind == "heading" and b.level == 1 for b in blocks):
        blocks = [Block("heading", c.title, 1), *blocks]
    props: dict[str, Any] = dict(c.meta)
    if c.title:
        props.setdefault("html_title", c.title)
    return blocks_to_document(blocks, "html", properties=props)


# --------------------------------------------------------------------------- docx


def parse_docx(data: bytes) -> ParsedDocument:
    from docx import Document as DocxDocument

    d = DocxDocument(io.BytesIO(data))
    blocks: list[Block] = []
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name if p.style is not None else "") or ""
        if style == "Title":
            blocks.append(Block("heading", text, 1))
        elif style.startswith("Heading"):
            m = re.search(r"(\d+)", style)
            blocks.append(Block("heading", text, int(m.group(1)) if m else 2))
        elif "List" in style:
            blocks.append(Block("bullet", text))
        else:
            blocks.append(Block("para", text))
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(Block("para", " | ".join(cells)))
    cp = d.core_properties
    props: dict[str, Any] = {
        "title": cp.title,
        "subject": cp.subject,
        "category": cp.category,
        "keywords": cp.keywords,
        "version": cp.version,
        "author": cp.author,
        "identifier": cp.identifier,
        "comments": cp.comments,
    }
    return blocks_to_document(blocks, "docx", properties={k: v for k, v in props.items() if v})


# --------------------------------------------------------------------------- pdf


def parse_pdf(data: bytes) -> ParsedDocument:
    import pdfplumber

    lines: list[tuple[str, float]] = []
    props: dict[str, Any] = {}
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        info = pdf.metadata or {}
        if info.get("Title"):
            props["title"] = info["Title"]
        for page in pdf.pages:
            words = page.extract_words(extra_attrs=["size"], keep_blank_chars=False, use_text_flow=True)
            current_top: float | None = None
            buf: list[str] = []
            sizes: list[float] = []
            for w in words:
                if current_top is None or abs(w["top"] - current_top) > 2.5:
                    if buf:
                        lines.append((" ".join(buf), max(sizes)))
                    buf, sizes, current_top = [], [], w["top"]
                buf.append(w["text"])
                sizes.append(round(float(w["size"]), 1))
            if buf:
                lines.append((" ".join(buf), max(sizes)))
    if not lines:
        raise IngestError("PDF has no extractable text (scanned PDFs need OCR, which is out of scope)")
    body_size = Counter(size for _, size in lines).most_common(1)[0][0]
    heading_sizes = sorted({size for _, size in lines if size > body_size + 0.5}, reverse=True)
    level_of = {size: min(i + 1, 3) for i, size in enumerate(heading_sizes)}
    blocks: list[Block] = []
    for text, size in lines:
        if size in level_of:
            blocks.append(Block("heading", text, level_of[size]))
            continue
        bullet = re.match(r"^[-•*]\s+(.*)$", text)
        if bullet:
            blocks.append(Block("bullet", bullet.group(1)))
            continue
        prev = blocks[-1] if blocks else None
        if (
            prev is not None
            and prev.kind in ("para", "bullet")
            and not re.search(r"[.:?!;]$", prev.text)
            and not KV_LINE.match(text)
            and not KV_LINE.match(prev.text)
        ):
            prev.text = f"{prev.text} {text}"
        else:
            blocks.append(Block("para", text))
    return blocks_to_document(blocks, "pdf", properties=props)


# --------------------------------------------------------------------------- dispatch

EXTENSIONS = {
    ".md": "md",
    ".markdown": "md",
    ".txt": "txt",
    ".html": "html",
    ".htm": "html",
    ".docx": "docx",
    ".pdf": "pdf",
}


def detect_format(filename: str) -> str:
    fmt = EXTENSIONS.get(Path(filename).suffix.lower())
    if fmt is None:
        raise IngestError(f"unsupported file type: {filename}")
    return fmt


def parse_bytes(filename: str, data: bytes) -> ParsedDocument:
    fmt = detect_format(filename)
    if fmt == "docx":
        return parse_docx(data)
    if fmt == "pdf":
        return parse_pdf(data)
    text = data.decode("utf-8", errors="replace").replace("\r\n", "\n")
    if fmt == "md":
        return parse_markdown(text)
    if fmt == "html":
        return parse_html(text)
    return parse_text(text)
